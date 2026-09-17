from pathlib import Path

import pandas as pd
import pytest
from docx import Document
from PIL import Image

from src.ingestion.models import UnsupportedFileTypeError
from src.ingestion.parsers import image as image_parser
from src.ingestion.parsers import pdf as pdf_parser
from src.ingestion.pipeline import FinancialIngestionPipeline


@pytest.fixture
def pipeline() -> FinancialIngestionPipeline:
    return FinancialIngestionPipeline()


def test_csv_parsing_preserves_table_and_metadata(pipeline: FinancialIngestionPipeline):
    result = pipeline.ingest(Path("data/samples/sample_pnl.csv"))

    assert result.total_chunks == 1
    chunk = result.chunks[0]
    assert chunk.chunk_type == "table"
    assert "| Revenue |" in chunk.content
    assert "$1200000" in chunk.content
    assert "$1450000" in chunk.content
    assert chunk.metadata["sheet_name"] == "sample_pnl"
    assert chunk.metadata["row_count"] == 3


def test_excel_parsing_preserves_each_sheet_and_metadata(
    pipeline: FinancialIngestionPipeline, tmp_path: Path
):
    workbook = tmp_path / "financials.xlsx"
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        pd.DataFrame({"Line Item": ["Revenue"], "Q1": [1200000]}).to_excel(
            writer, sheet_name="P&L", index=False
        )
        pd.DataFrame({"Account": ["Cash"], "Balance": [500000]}).to_excel(
            writer, sheet_name="Balance Sheet", index=False
        )

    result = pipeline.ingest(workbook)

    assert result.total_chunks == 2
    sheets = {chunk.metadata["sheet_name"]: chunk for chunk in result.chunks}
    assert "| Revenue | 1200000 |" in sheets["P&L"].content
    assert "| Cash | 500000 |" in sheets["Balance Sheet"].content
    assert all(chunk.chunk_type == "table" for chunk in result.chunks)


def test_text_parsing_creates_searchable_chunk(
    pipeline: FinancialIngestionPipeline, tmp_path: Path
):
    text_path = tmp_path / "notes.txt"
    text_path.write_text("Revenue increased in Q2.", encoding="utf-8")

    result = pipeline.ingest(text_path)

    assert result.total_chunks == 1
    assert result.chunks[0].content == "Revenue increased in Q2."
    assert result.chunks[0].chunk_type == "text"


def test_docx_parsing_extracts_non_empty_paragraphs(
    pipeline: FinancialIngestionPipeline, tmp_path: Path
):
    document_path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("Revenue increased in Q2.")
    document.add_paragraph("")
    document.add_paragraph("Cash remained stable.")
    document.save(document_path)

    result = pipeline.ingest(document_path)

    assert result.total_chunks == 1
    assert result.chunks[0].content == "Revenue increased in Q2.\nCash remained stable."


def test_image_parsing_extracts_ocr_text(
    pipeline: FinancialIngestionPipeline, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    image_path = tmp_path / "receipt.png"
    Image.new("RGB", (120, 40), "white").save(image_path)
    monkeypatch.setattr(
        image_parser.pytesseract,
        "image_to_string",
        lambda image: "Total $125.00",
    )

    result = pipeline.ingest(image_path)

    assert result.total_chunks == 1
    assert result.chunks[0].content == "Total $125.00"
    assert result.chunks[0].chunk_type == "image_ocr"
    assert result.chunks[0].metadata["ocr"] is True


def test_invalid_extension_is_rejected(pipeline: FinancialIngestionPipeline, tmp_path: Path):
    invalid_file = tmp_path / "financials.rtf"
    invalid_file.write_text("not a supported document", encoding="utf-8")

    with pytest.raises(UnsupportedFileTypeError):
        pipeline.ingest(invalid_file)


def test_pdf_table_extraction_is_markdown_and_page_aware(
    pipeline: FinancialIngestionPipeline, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    pdf_path = tmp_path / "report.pdf"

    class FakePage:
        def extract_tables(self):
            return [[["Line Item", "Q1"], ["Revenue", "$1200000"]]]

        def extract_text(self):
            return ""

    class FakePDF:
        pages = [FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc_value, traceback):
            return False

    monkeypatch.setattr(pdf_parser.pdfplumber, "open", lambda path: FakePDF())

    result = pipeline.ingest(pdf_path)

    assert result.total_chunks == 1
    chunk = result.chunks[0]
    assert chunk.chunk_type == "table"
    assert "| Revenue | $1200000 |" in chunk.content
    assert chunk.metadata["page_number"] == 1